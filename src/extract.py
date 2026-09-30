"""Text extraction with a refusal rule: if too little text comes out (e.g. a scanned PDF), we refuse to score
rather than return a confident NO_FLAG on an empty document."""
import io, re
MIN_WORDS = 150

def extract_text(name, data):
    n = name.lower()
    if n.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if n.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        parts += [c.text for t in d.tables for row in t.rows for c in row.cells]
        return "\n".join(parts)
    return data.decode("utf-8", errors="replace")

def check_text(text):
    """Returns (ok, message)."""
    words = re.findall(r"[A-Za-z]{2,}", text or "")
    if len(words) < MIN_WORDS:
        return False, (f"Only {len(words)} words could be read (need at least {MIN_WORDS}). "
                       "If this is a scanned PDF, run OCR first or paste the text. Nothing was scored.")
    return True, f"{len(words):,} words read."
